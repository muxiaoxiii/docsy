"""Isolated native converter comparison using synthetic documents, never Office automation."""
from pathlib import Path
import concurrent.futures
import hashlib
import json
import subprocess
import sys
import time

from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

repo = Path.cwd()
root = repo / "release-candidate/office-converter-upgrade-2026-10-19"
upstream = Path(Path("/tmp/docsy-office2pdf-current-path").read_text())
for name in ["inputs", "office2pdf", "minipdf", "renders"]:
    (root / name).mkdir(parents=True, exist_ok=True)

def base():
    d = Document()
    section = d.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.8)
    normal = d.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Songti SC")
    return d

d = base()
d.sections[0].header.paragraphs[0].text = "AUDIT HEADER"
d.sections[0].footer.paragraphs[0].text = "AUDIT FOOTER"
d.add_heading("Native conversion probe", 0)
d.add_paragraph("BASIC FIRST PAGE 合同中文测试。")
p = d.add_paragraph("Formatting: ")
p.add_run("BOLD ").bold = True
p.add_run("ITALIC ").italic = True
d.add_page_break()
d.add_paragraph("BASIC SECOND PAGE 第二页内容。")
d.save(root / "inputs/basic.docx")

d = base()
d.add_heading("Table conversion probe", 0)
t = d.add_table(rows=4, cols=3)
t.style = "Table Grid"
for row, vals in zip(t.rows, [["ITEM", "DESCRIPTION", "VALUE"], ["A", "中文表格内容", "100"], ["B", "Two line\ncell content", "200"], ["C", "MERGED CELL", "300"]]):
    for cell, val in zip(row.cells, vals):
        cell.text = val
t.cell(3, 1).merge(t.cell(3, 2)).text = "MERGED CELL"
d.add_paragraph("TABLE END MARKER")
d.save(root / "inputs/table.docx")

d = base()
d.add_heading("Section conversion probe", 0)
d.add_paragraph("PORTRAIT SECTION")
section = d.add_section(WD_SECTION_START.NEW_PAGE)
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = Inches(11), Inches(8.5)
d.add_paragraph("LANDSCAPE SECTION 横向页面")
d.save(root / "inputs/sections.docx")

sidecar = subprocess.run([str(repo/"src-tauri/runtime/bin/doc2x"), str(repo/"src-tauri/tests/fixtures/doc-import/sample.doc"), "-o", str(root/"inputs/legacy.docx"), "-v", "error"], capture_output=True, text=True, timeout=30)
assert sidecar.returncode == 0 and (root/"inputs/legacy.docx").is_file(), sidecar.stderr

engines = {
    "office2pdf": upstream/"office2pdf-v0.9.1-aarch64-apple-darwin/office2pdf",
    "minipdf": upstream/"MiniPdf/minipdf-rs/target/debug/minipdf",
}

def convert(job):
    engine, source = job
    output = root/engine/source.with_suffix(".pdf").name
    started = time.perf_counter()
    p = subprocess.run([str(engines[engine]),str(source),"-o",str(output)],capture_output=True,text=True,timeout=60)
    return {"engine":engine,"input":source.name,"exit_code":p.returncode,"elapsed_s":round(time.perf_counter()-started,3),"output_exists":output.is_file(),"output_bytes":output.stat().st_size if output.is_file() else 0,"stderr":p.stderr,"stdout":p.stdout}

jobs = [(engine, source) for engine in engines for source in sorted((root/"inputs").glob("*.docx"))]
with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
    results = list(pool.map(convert, jobs))

from pypdf import PdfReader
from PIL import Image, ImageDraw, ImageOps
tiles = []
for r in results:
    if not r["output_exists"]:
        continue
    pdf = root/r["engine"]/Path(r["input"]).with_suffix(".pdf")
    reader = PdfReader(pdf)
    r["pages"] = len(reader.pages)
    r["page_boxes"] = [[float(v) for v in page.mediabox] for page in reader.pages]
    r["page_text"] = [page.extract_text() for page in reader.pages]
    prefix = root/"renders"/f"{Path(r['input']).stem}-{r['engine']}"
    render = subprocess.run(["pdftoppm","-r","80","-png",str(pdf),str(prefix)],capture_output=True,text=True,timeout=40)
    r["render_code"] = render.returncode
    for image in sorted((root/"renders").glob(prefix.name+"-*.png")):
        im = Image.open(image).convert("RGB")
        tile = Image.new("RGB",(330,455),"#eeeeee")
        ImageDraw.Draw(tile).text((6,5),image.stem,fill="black")
        tile.paste(ImageOps.contain(im,(318,420)),(6,28))
        tiles.append(tile)

sheet = Image.new("RGB",(330*4,455*((len(tiles)+3)//4)),"white")
for i,tile in enumerate(tiles):sheet.paste(tile,((i%4)*330,(i//4)*455))
sheet.save(root/"contact-sheet.png")

out = {"scope":"Synthetic smoke test; no native Office ground truth; debug MiniPdf and release office2pdf timings are not comparable benchmarks", "office2pdf_version":"0.9.1", "minipdf_source":subprocess.check_output(["git","-C",str(upstream/"MiniPdf"),"rev-parse","HEAD"],text=True).strip(),"legacy_docx_sidecar_exit":sidecar.returncode,"results":results}
(root/"results.json").write_text(json.dumps(out,ensure_ascii=False,indent=2)+"\n")
for r in results:print(r["engine"],r["input"],r["exit_code"],r.get("pages"),r.get("page_boxes"))
